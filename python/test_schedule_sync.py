import io
import sys
from email.message import EmailMessage
import unittest
from datetime import datetime
from pathlib import Path
from unittest.mock import MagicMock, patch

import openpyxl
import schedule_sync as sync
import import_committed_exams as committed
import subprocess
import tempfile


class ScheduleSyncTests(unittest.TestCase):
    def test_push_range_handles_spaces_multiple_commits_and_deletions(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            def git(*args):
                return subprocess.check_output(["git", "-C", directory, *args]).decode().strip()
            git("init", "-q")
            git("config", "user.email", "test@example.com")
            git("config", "user.name", "Test")
            (root / "deleted.xlsx").write_bytes(b"old")
            git("add", ".")
            git("commit", "-qm", "base")
            before = git("rev-parse", "HEAD")
            (root / "deleted.xlsx").unlink()
            (root / "Exam schedule with spaces.xlsx").write_bytes(self.workbook())
            git("add", "-A")
            git("commit", "-qm", "schedule")
            (root / "exam-schedules").mkdir()
            (root / "exam-schedules" / "engineering.xlsx").write_bytes(self.workbook(("FSE",)))
            (root / "~$temporary.xlsx").write_bytes(b"lock")
            git("add", ".")
            git("commit", "-qm", "second school")
            after = git("rev-parse", "HEAD")
            with patch.object(sync, "ROOT", root):
                expected = ["Exam schedule with spaces.xlsx", "exam-schedules/engineering.xlsx"]
                self.assertEqual(committed.changed_workbooks(before, after), expected)
                self.assertEqual(committed.changed_workbooks("0" * 40, after), expected)
                with patch.object(sys, "argv", ["import_committed_exams.py", "--before", before, "--after", after]), \
                     patch.object(sync, "connect") as connect:
                    committed.main()
                    connect.assert_not_called()

    def workbook(self, names=("FSC",)):
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        for name in names:
            ws = wb.create_sheet(name)
            ws.append(["Date", "9:00 to 10:00 AM", None, "10:20 to 11:20 AM"])
            ws.append([datetime(2026, 9, 19), "CS1001\nProgramming\nBS(CS) A,B\n2026"])
        out = io.BytesIO()
        wb.save(out)
        return out.getvalue()

    def test_scopes_and_course_continuation(self):
        row = sync.parser.parse_exam_paper_cell("EE2003\nComputer Organization\nBCE-A\nBEE-A,B\n2025")
        self.assertEqual(row["course"], "Computer Organization")
        self.assertEqual(row["sections"], {"CE": ["A"], "EE": ["A", "B"]})
        self.assertEqual(sync.parser.parse_exam_paper_cell("AF1001 Accounting\nBBA-A,B\n2026")["sections"], {"BBA": ["A", "B"]})

    def test_multiple_sheets_accumulate(self):
        docs = sync.parse_attachment(self.workbook(("FSC A", "FSC B")), "exam.xlsx", "exams")
        self.assertEqual(docs["exams/computing"]["count"], 2)

    def test_room_grid_merges_and_four_slots(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "FSC"
        ws.append(["Date", "Room", "9:00 to 10:00", None, "10:20 to 11:20", None,
                   "11:40 to 12:40", None, "13:00 to 14:00"])
        ws.append([datetime(2026, 9, 19), "C-301", "CS-1001 Programming", None, None, None, None, None, "MT-1001 Calculus"])
        ws.append([None, "C-302"])
        ws.merge_cells("A2:A3")
        ws.merge_cells("C2:C3")
        out = io.BytesIO()
        wb.save(out)
        doc = sync.parse_attachment(out.getvalue(), "exam.xlsx", "exams")["exams/computing"]
        self.assertEqual(len(doc["flat_exams"]), 3)
        self.assertFalse(doc["exams"])
        self.assertEqual(doc["flat_exams"][-1]["room"], "C-302")

    def test_routes(self):
        self.assertEqual(sync.route("Show Up Schedule", "schedule.xlsx"), "showup")
        self.assertEqual(sync.route("Schedule of Final Examination", "attachment.xlsx"), "exams")
        self.assertEqual(sync.route("Seating Plan", "plan.pdf"), "seating")
        self.assertIsNone(sync.route("Class schedule", "classes.xlsx"))

    def test_invalid_records_rejected_before_connection(self):
        with self.assertRaises(ValueError):
            sync.publish(None, {"exams/computing": {"count": 0, "exams": []}}, {})
        docs = sync.parse_attachment(self.workbook(), "exam.xlsx", "exams")
        docs["exams/computing"]["exams"][0]["date"] = None
        with self.assertRaises(ValueError):
            sync.validate(docs)

    def test_incomplete_pdf_rejected(self):
        with patch.object(sync.parser, "parse_seating_plan_email", return_value={"room_occupancy": {"complete": False}}):
            with self.assertRaisesRegex(ValueError, "Incomplete seating"):
                sync.parse_attachment(b"pdf", "seating.pdf", "seating")

    def database(self, old):
        db = MagicMock()
        db.documents.find_one.return_value = old
        session = db.client.start_session.return_value.__enter__.return_value
        session.with_transaction.side_effect = lambda fn, **kwargs: fn(session)
        return db

    def test_unchanged_import_has_no_document_write(self):
        docs = sync.parse_attachment(self.workbook(), "exam.xlsx", "exams")
        db = self.database({"data": docs["exams/computing"]})
        result = sync.publish(db, docs, {"received_at": "2026-09-11T00:00:00+00:00"}, "uid")
        self.assertEqual(result["exams/computing"], "unchanged")
        db.documents.replace_one.assert_not_called()
        db.schedule_imports.update_one.assert_called_once()

    def test_older_revision_cannot_replace_newer(self):
        docs = sync.parse_attachment(self.workbook(), "exam.xlsx", "exams")
        old = {"data": {**docs["exams/computing"], "extra": "new"},
               "source": {"received_at": "2026-09-12T00:00:00+00:00"}}
        db = self.database(old)
        result = sync.publish(db, docs, {"received_at": "2026-09-11T00:00:00+00:00"})
        self.assertIn("older", result["exams/computing"])
        db.documents.replace_one.assert_not_called()

    def test_write_archives_previous_document_in_same_session(self):
        docs = sync.parse_attachment(self.workbook(), "exam.xlsx", "exams")
        db = self.database({"_id": "exams/computing", "data": {**docs["exams/computing"], "extra": "old"}})
        result = sync.publish(db, docs, {"received_at": "2026-09-11T00:00:00+00:00"}, "uid")
        self.assertEqual(result["exams/computing"], "written")
        session = db.documents.replace_one.call_args.kwargs["session"]
        self.assertIs(db.schedule_revisions.update_one.call_args.kwargs["session"], session)
        self.assertIs(db.schedule_imports.update_one.call_args.kwargs["session"], session)

    def test_supplied_workbook(self):
        files = list(Path(sync.ROOT).glob("1st Sessional*11-09-2026.xlsx"))
        if not files:
            self.skipTest("Source workbook not present")
        docs = sync.parse_attachment(files[0].read_bytes(), files[0].name, "exams")
        self.assertEqual({k: d["count"] for k, d in docs.items()},
                         {"exams/computing": 109, "exams/business": 98, "exams/engineering": 45})
        self.assertEqual(len({e["time"] for e in docs["exams/computing"]["exams"]}), 8)
        self.assertTrue(all(e["course"] for d in docs.values() for e in d["exams"]))

    def test_gmail_failure_retries_without_marking_read_and_continues(self):
        msg = EmailMessage()
        msg["Subject"] = "Sessional Exam Schedule"
        msg.set_content("Attached")
        msg.add_attachment(self.workbook(), maintype="application", subtype="octet-stream", filename="exam.xlsx")
        mail = MagicMock()
        mail.select.return_value = ("OK", [])
        mail.response.return_value = ("UIDVALIDITY", [b"123"])
        def uid(command, *args):
            if command == "search":
                return "OK", [b"1 2"]
            if command == "fetch":
                self.assertIn("BODY.PEEK[]", args[1])
                return "OK", [(b'1 (INTERNALDATE "11-Sep-2026 12:00:00 +0000")', msg.as_bytes())]
            return "OK", []
        mail.uid.side_effect = uid
        db = MagicMock()
        db.schedule_imports.find_one.return_value = None
        with patch.dict(sync.os.environ, {"GMAIL_USER": "test@example.com", "GMAIL_PASS": "fixture"}), \
             patch.object(sync, "connect", return_value=db), \
             patch("imaplib.IMAP4_SSL", return_value=mail), \
             patch.object(sync, "publish", side_effect=[RuntimeError("write failed"), {"exams/computing": "written"}]) as publish:
            with self.assertRaisesRegex(RuntimeError, "1 message"):
                sync.sync_gmail()
        self.assertEqual(publish.call_count, 2)
        stores = [c for c in mail.uid.call_args_list if c.args[0] == "store"]
        self.assertEqual(len(stores), 1)
        self.assertEqual(stores[0].args[1], b"2")


if __name__ == "__main__":
    unittest.main()
